"""Extract full text (with tables) from collected 4-가 attachments on macOS/Linux.

The original extraction ran on Windows through ``kordoc-fixed.cmd``.  This
rewrite produces rows in the same shape so every downstream step
(``build_biology_extraction_retry_queue`` -> ``build_biology_assessment_evidence_index``
-> catalogue -> publish DB) runs unchanged::

    {"source": {...manifest row...}, "sha256": "...", "index_status": null,
     "extract_status": "ok|short|failed|missing", "char_count": 12345,
     "subject_hits": [{"subject": "...", "curriculum_hint": "..."}],
     "curriculum_code_hints": ["2015", "2022"], "text": "..."}

``text`` keeps the converter conventions the detail parser expects: plain
paragraph lines, ``### `` headings for numbered plan headings, tables as
sanitised ``<table>`` HTML with rowspan/colspan (cell paragraphs joined with
``<br>``), and one ``# <member name>`` heading per ZIP member.

Converters:
  .hwp   pyhwp ``hwp5html`` (XHTML -> tables/paragraphs); slow (~15 s/file)
  .hwpx  OWPML section XML (zip) parsed directly
  .pdf   pdfplumber page text + detected tables
  .docx  python-docx paragraphs and tables
  .xlsx  openpyxl sheets as tables
  .zip   members extracted and converted with the rules above

Usage:
  python scripts/extract_biology_candidate_text.py --manifest data/derived/biology_assessment_source_manifest.csv \\
      --output data/derived/biology_allplan_remaining_text_0000.jsonl --offset 0 --limit 500 --workers 4
  python scripts/extract_biology_candidate_text.py --input-jsonl data/derived/biology_assessment_retry_queue_v2.jsonl \\
      --output data/derived/biology_assessment_retry_fulltext_v2.jsonl
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html as html_lib
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from concurrent.futures import ProcessPoolExecutor
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from biology_subject_aliases import subject_hits  # noqa: E402

SHORT_TEXT_CHARS = 300
CODE_2015_RE = re.compile(r"\[(?:10|12)[가-힣ⅠⅡ]{1,8}\d{2}-\d{2}\]")
CODE_2022_RE = re.compile(r"\[(?:10|12)[가-힣]{1,8}[12]-\d{2}-\d{2}\]")
HEADING_RE = re.compile(r"^\s*(?:\d+\.|[가-힣]\.|[ⅠⅡⅢⅣⅤⅥⅦⅧ]\.|\(\d+\)|[①-⑳])\s*\S")
HWPX_NS = {"hp": "http://www.hancom.co.kr/hwpml/2011/paragraph"}


# ---------------------------------------------------------------- helpers ---
def curriculum_code_hints(text: str) -> list[str]:
    hints = []
    if CODE_2015_RE.search(text):
        hints.append("2015")
    if CODE_2022_RE.search(text):
        hints.append("2022")
    return hints


def escape(value: str) -> str:
    return html_lib.escape(value, quote=False)


def table_html(rows: list[list[dict]]) -> str:
    """Render rows of {text, rowspan, colspan, nested} cells as a minimal safe table.

    A one-cell table is a page frame or a boxed heading, not data: its content
    is emitted at the top level so the assessment tables inside it are not
    buried in a cell.  ``nested`` holds already-rendered inner tables.
    """

    if len(rows) == 1 and len(rows[0]) == 1:
        cell = rows[0][0]
        parts = [paragraph_line(line) for line in str(cell.get("text") or "").split("\n")]
        parts = [part for part in parts if part]
        parts.extend(cell.get("nested") or [])
        return "\n".join(parts)
    out = ["<table>"]
    for row in rows:
        cells = []
        for cell in row:
            attrs = ""
            if int(cell.get("rowspan") or 1) > 1:
                attrs += f' rowspan="{int(cell["rowspan"])}"'
            if int(cell.get("colspan") or 1) > 1:
                attrs += f' colspan="{int(cell["colspan"])}"'
            body = escape(str(cell.get("text") or "")).replace("\n", "<br>")
            nested = "".join(cell.get("nested") or [])
            cells.append(f"<td{attrs}>{body}{nested}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def paragraph_line(text: str) -> str:
    text = re.sub(r"[ \t ]+", " ", text).strip()
    if not text:
        return ""
    if HEADING_RE.match(text) and len(text) <= 80:
        return f"### {text}"
    return text


# ------------------------------------------------------------------ HWPX ---
HP_P = f"{{{HWPX_NS['hp']}}}p"
HP_T = f"{{{HWPX_NS['hp']}}}t"
HP_TBL = f"{{{HWPX_NS['hp']}}}tbl"
HP_TR = f"{{{HWPX_NS['hp']}}}tr"
HP_TC = f"{{{HWPX_NS['hp']}}}tc"
HP_SPAN = f"{{{HWPX_NS['hp']}}}cellSpan"


def _has_ancestor(parent: dict, node, tag: str, stop=None) -> bool:
    current = parent.get(node)
    while current is not None and current is not stop:
        if current.tag == tag:
            return True
        current = parent.get(current)
    return False


def hwpx_cell_text(parent: dict, tc) -> str:
    """Cell text: paragraphs directly in the cell, excluding nested tables."""

    parts = []
    for paragraph in tc.iter(HP_P):
        if _has_ancestor(parent, paragraph, HP_TBL, stop=tc):
            continue
        line = "".join(t.text or "" for t in paragraph.iter(HP_T)).strip()
        if line:
            parts.append(line)
    return "\n".join(parts)


def hwpx_table(parent: dict, tbl) -> str:
    rows = []
    for tr in tbl.findall(HP_TR):
        cells = []
        for tc in tr.findall(HP_TC):
            span = tc.find(HP_SPAN)
            nested = [
                hwpx_table(parent, inner)
                for inner in tc.iter(HP_TBL)
                if not _has_ancestor(parent, inner, HP_TBL, stop=tc)
            ]
            cells.append({
                "text": hwpx_cell_text(parent, tc),
                "nested": [html for html in nested if html],
                "rowspan": span.get("rowSpan") if span is not None else 1,
                "colspan": span.get("colSpan") if span is not None else 1,
            })
        rows.append(cells)
    return table_html(rows) if rows else ""


def extract_hwpx(data: bytes) -> str:
    lines: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sections = sorted(
            name for name in archive.namelist() if re.search(r"Contents/section\d+\.xml$", name)
        )
        for name in sections:
            root = ElementTree.fromstring(archive.read(name))
            parent = {child: node for node in root.iter() for child in node}
            for element in root:
                if element.tag != HP_P:
                    continue
                tables = [
                    tbl for tbl in element.iter(HP_TBL)
                    if not _has_ancestor(parent, tbl, HP_TBL, stop=element)
                ]
                text = "".join(
                    t.text or "" for t in element.iter(HP_T)
                    if not _has_ancestor(parent, t, HP_TBL, stop=element)
                )
                line = paragraph_line(text)
                if line:
                    lines.append(line)
                for tbl in tables:
                    lines.append(hwpx_table(parent, tbl))
                    lines.append("")
    return "\n".join(lines)


# ------------------------------------------------------------------- HWP ---
class _XhtmlToText(HTMLParser):
    """Walk pyhwp's XHTML and emit paragraphs and nested-aware tables."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self.table_stack: list[list[list[dict]]] = []
        self.cell_stack: list[dict] = []
        self.text: list[str] = []

    def _flush_paragraph(self) -> None:
        text = "".join(self.text)
        self.text = []
        if self.cell_stack:
            cell = self.cell_stack[-1]
            if text.strip():
                cell["text"] = (cell["text"] + "\n" if cell["text"] else "") + text.strip()
            return
        line = paragraph_line(text)
        if line:
            self.lines.append(line)

    def handle_starttag(self, tag: str, attrs) -> None:
        attributes = dict(attrs)
        if tag == "table":
            self._flush_paragraph()
            self.table_stack.append([])
        elif tag == "tr" and self.table_stack:
            self.table_stack[-1].append([])
        elif tag in ("td", "th") and self.table_stack:
            cell = {"text": "", "nested": [], "rowspan": attributes.get("rowspan") or 1,
                    "colspan": attributes.get("colspan") or 1}
            self.table_stack[-1][-1].append(cell)
            self.cell_stack.append(cell)
        elif tag == "br":
            self.text.append("\n")
        elif tag == "p" and not self.cell_stack:
            self._flush_paragraph()

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self.cell_stack:
            self._flush_paragraph()
            self.cell_stack.pop()
        elif tag == "table" and self.table_stack:
            rows = self.table_stack.pop()
            rendered = table_html(rows) if rows else ""
            if self.cell_stack:
                if rendered:
                    self.cell_stack[-1]["nested"].append(rendered)
            elif rendered:
                self.lines.append(rendered)
                self.lines.append("")
        elif tag == "p" and not self.cell_stack:
            self._flush_paragraph()

    def handle_data(self, data: str) -> None:
        self.text.append(data.replace("\r", ""))


def extract_hwp(data: bytes) -> str:
    hwp5html = shutil.which("hwp5html") or str(Path(sys.executable).with_name("hwp5html"))
    with tempfile.TemporaryDirectory(prefix="hwp5_") as workdir:
        source = Path(workdir) / "input.hwp"
        source.write_bytes(data)
        output = Path(workdir) / "out"
        subprocess.run(
            [hwp5html, "--output", str(output), str(source)],
            check=True, capture_output=True, timeout=600,
        )
        xhtml = (output / "index.xhtml").read_text(encoding="utf-8", errors="replace")
    xhtml = re.sub(r"<style\b.*?</style>", "", xhtml, flags=re.S | re.I)
    parser = _XhtmlToText()
    parser.feed(xhtml)
    parser.close()
    parser._flush_paragraph()
    return "\n".join(parser.lines)


# ------------------------------------------------------------------- PDF ---
def extract_pdf(data: bytes) -> str:
    import pdfplumber

    lines: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for raw in text.splitlines():
                line = paragraph_line(raw)
                if line:
                    lines.append(line)
            for table in page.extract_tables() or []:
                rows = [[{"text": (cell or "").strip()} for cell in row] for row in table if row]
                if rows:
                    lines.append(table_html(rows))
                    lines.append("")
    return "\n".join(lines)


# ------------------------------------------------------------ DOCX/XLSX ---
def extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        line = paragraph_line(paragraph.text)
        if line:
            lines.append(line)
    for table in document.tables:
        rows = [[{"text": cell.text.strip()} for cell in row.cells] for row in table.rows]
        lines.append(table_html(rows))
        lines.append("")
    return "\n".join(lines)


def extract_xlsx(data: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"### {sheet.title}")
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if row is None or all(value in (None, "") for value in row):
                continue
            rows.append([{"text": "" if value is None else str(value)} for value in row])
        if rows:
            lines.append(table_html(rows))
        lines.append("")
    return "\n".join(lines)


# ------------------------------------------------------------- dispatch ---
def sniff_extension(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower().lstrip(".")
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        if suffix in ("hwpx", "docx", "xlsx", "zip"):
            return suffix
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            return suffix or "unknown"
        if any(n.startswith("Contents/section") for n in names):
            return "hwpx"
        if any(n.startswith("word/") for n in names):
            return "docx"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        return "zip"
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return "hwp" if suffix in ("hwp", "", "unknown") else suffix
    return suffix or "unknown"


def extract_bytes(name: str, data: bytes, depth: int = 0) -> str:
    kind = sniff_extension(name, data)
    if kind == "hwp":
        return extract_hwp(data)
    if kind == "hwpx":
        return extract_hwpx(data)
    if kind == "pdf":
        return extract_pdf(data)
    if kind == "docx":
        return extract_docx(data)
    if kind == "xlsx":
        return extract_xlsx(data)
    if kind == "zip" and depth < 2:
        parts: list[str] = []
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for info in archive.infolist():
                if info.is_dir():
                    continue
                member_name = decode_zip_name(info)
                if member_name.startswith("__MACOSX") or Path(member_name).name.startswith("."):
                    continue
                member = archive.read(info)
                try:
                    text = extract_bytes(member_name, member, depth + 1)
                except Exception as exc:  # noqa: BLE001 - one bad member must not sink the archive
                    text = f"[extract_failed: {type(exc).__name__}]"
                parts.append(f"# {Path(member_name).name}\n\n{text}")
        return "\n\n".join(parts)
    raise ValueError(f"unsupported document type: {kind}")


def decode_zip_name(info: zipfile.ZipInfo) -> str:
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp949")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return info.filename


def extract_row(task: tuple[dict, str]) -> dict:
    source, project_root = task
    saved = str(source.get("saved_path") or "")
    path = Path(saved)
    if not path.is_absolute():
        path = Path(project_root) / path
    row = {
        "source": source, "sha256": "", "index_status": None, "extract_status": "missing",
        "char_count": 0, "subject_hits": [], "curriculum_code_hints": [], "text": "",
    }
    if not path.is_file():
        return row
    data = path.read_bytes()
    row["sha256"] = hashlib.sha256(data).hexdigest()
    try:
        text = extract_bytes(path.name, data)
    except Exception as exc:  # noqa: BLE001 - recorded as a failed row, not a crash
        row["extract_status"] = "failed"
        row["text"] = f"[extract_failed: {type(exc).__name__}: {exc}]"[:300]
        return row
    text = text.strip()
    row.update({
        "extract_status": "ok" if len(text) >= SHORT_TEXT_CHARS else "short",
        "char_count": len(text),
        "subject_hits": subject_hits(text),
        "curriculum_code_hints": curriculum_code_hints(text),
        "text": text,
    })
    return row


def load_sources(args) -> list[dict]:
    sources: list[dict] = []
    if args.manifest:
        with args.manifest.open("r", encoding="utf-8-sig", newline="") as handle:
            sources.extend(dict(row) for row in csv.DictReader(handle))
    for path in args.input_jsonl:
        with path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if line.strip():
                    record = json.loads(line)
                    sources.append(record.get("source") or record)
    return sources


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--manifest", type=Path, default=None, help="source manifest CSV")
    parser.add_argument("--input-jsonl", type=Path, action="append", default=[],
                        help="rows carrying a source dict (retry queue); repeatable")
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--kordoc", type=Path, default=None, help="ignored (Windows converter)")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--workers", type=int, default=2)
    args = parser.parse_args()

    sources = load_sources(args)
    if not sources:
        raise SystemExit("no source rows: pass --manifest or --input-jsonl")
    selected = sources[args.offset:]
    if args.limit:
        selected = selected[: args.limit]
    tasks = [(source, str(args.project_root)) for source in selected]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    statuses: dict[str, int] = {}
    with args.output.open("w", encoding="utf-8", newline="\n") as out:
        if args.workers > 1:
            iterator = ProcessPoolExecutor(max_workers=args.workers).map(extract_row, tasks)
        else:
            iterator = map(extract_row, tasks)
        for index, row in enumerate(iterator, 1):
            out.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
            statuses[row["extract_status"]] = statuses.get(row["extract_status"], 0) + 1
            if index % 50 == 0:
                print(f"progress {index}/{len(tasks)} {statuses}", flush=True)
    print(json.dumps({"rows": len(tasks), "extract_status": statuses}, ensure_ascii=False))


if __name__ == "__main__":
    main()
